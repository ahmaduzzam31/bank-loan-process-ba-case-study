"""Build deliverables/07_business_case.xlsx: options appraisal with NPV, IRR, payback and sensitivity.

Volumes come from the process analysis (reports/opportunity.json). Every calculation is an Excel formula.
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables/07_business_case.xlsx"
opp = json.loads((ROOT / "reports/opportunity.json").read_text())

FONT = "Arial"
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
FIN = Font(name=FONT, size=10, color="0000FF")
FH = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FT = Font(name=FONT, size=15, bold=True, color="1F3864")
FS = Font(name=FONT, size=9, italic=True, color="52514E")
FILL_H = PatternFill("solid", fgColor="1F3864")
FILL_IN = PatternFill("solid", fgColor="FFFF00")
FILL_SEC = PatternFill("solid", fgColor="DCE6F2")
THIN = Side(style="thin", color="C3C2B7")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
EUR = '€#,##0;(€#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'

wb = Workbook()

# ---------------------------------------------------------------- Inputs
wi = wb.active
wi.title = "Inputs"
wi["A1"], wi["A2"] = "Business Case Inputs: SME Loan Origination Redesign", \
    "Harbour Bank Ireland (fictional). Blue = input; yellow = key assumption. Source column documents every number."
wi["A1"].font, wi["A2"].font = FT, FS
wi.sheet_view.showGridLines = False
for c, w in zip("ABCDE", (46, 14, 14, 14, 70)):
    wi.column_dimensions[c].width = w
for j, h in enumerate(["Input", "A. Do nothing", "B. Quick wins", "C. Digital + quick wins", "Source / rationale"], 1):
    cell = wi.cell(row=4, column=j, value=h)
    cell.font, cell.fill, cell.border = FH, FILL_H, BOX
    cell.alignment = Alignment(wrap_text=True, horizontal="center")

rows = [
    ("GENERAL", None, None, None, None, None),
    ("Discount rate (annual)", 0.08, None, None, PCT, "Bank hurdle rate for change projects (assumption)."),
    ("Average funded loan (€)", opp["avg_funded_loan_eur"], None, None, EUR, "Process analysis: mean amount of funded applications."),
    ("Net margin per year on a funded loan", 0.018, None, None, PCT, "After funding cost, expected credit loss and servicing (assumption)."),
    ("Average loan life (years)", 3, None, None, "0.0", "SME term loans 1 to 7 years; average outstanding life (assumption)."),
    ("Benefit realisation factor", 0.5, None, None, PCT, "Only half of the modelled extra loans are counted, for prudence."),
    ("Underwriter cost per hour (€)", 55, None, None, EUR, "Fully loaded cost (assumption)."),
    ("Relationship manager cost per hour (€)", 45, None, None, EUR, "Fully loaded cost (assumption)."),
    ("VOLUMES (per year)", None, None, None, None, None),
    ("Extra funded loans vs today (modelled)", 0, opp["extra_funded_loans_option_b"], opp["extra_funded_loans_option_c"], "#,##0",
     "Process analysis notebook, section 5: fewer withdrawals from faster decisions and complete applications."),
    ("Credit assessments", opp["assessments_per_year"], opp["assessments_per_year"], opp["assessments_per_year"], "#,##0", "Event log, last 12 months."),
    ("Document requests (today)", opp["document_requests_per_year"], opp["document_requests_per_year"],
     opp["document_requests_per_year"], "#,##0", "Event log, last 12 months."),
    ("EFFICIENCY", None, None, None, None, None),
    ("Minutes re-keying financials per assessment", 45, 45, 45, "0", "Time study with underwriters (workshop estimate)."),
    ("Reduction in re-keying", 0, 0, 0.8, PCT, "C: open banking / accounting-software data feed replaces manual spreading."),
    ("Minutes of RM time per document request", 40, 40, 40, "0", "RM workshop estimate (call, email, chase, check)."),
    ("Reduction in document requests", 0, 0.5, 0.8, PCT, "B: digital checklist & upload. C: plus pre-filled data via open banking."),
    ("COSTS (€)", None, None, None, None, None),
    ("One-off cost: year 0", 0, 450000, 1500000, EUR, "B: policy change, checklist, portal config, training. C: adds platform build."),
    ("One-off cost: year 1", 0, 0, 900000, EUR, "C: integration and migration in year 1."),
    ("Annual run cost (from year 1)", 0, 120000, 380000, EUR, "Licences, support, credit forum time."),
    ("RAMP-UP of benefits", None, None, None, None, None),
    ("Year 1", 0, 0.5, 0.25, PCT, "Share of full benefit achieved."),
    ("Year 2", 0, 1, 0.75, PCT, ""),
    ("Year 3", 0, 1, 1, PCT, ""),
    ("Year 4", 0, 1, 1, PCT, ""),
    ("Year 5", 0, 1, 1, PCT, ""),
]
ref = {}
for i, (label, a, b, c, fmt, note) in enumerate(rows, start=5):
    lab = wi.cell(row=i, column=1, value=label)
    if a is None and b is None:
        lab.font, lab.fill = FB, FILL_SEC
        for j in range(2, 6):
            wi.cell(row=i, column=j).fill = FILL_SEC
        continue
    lab.font = F
    vals = [a, b, c] if b is not None else [a]
    for j, v in enumerate(vals, start=2):
        cell = wi.cell(row=i, column=j, value=v)
        cell.font, cell.number_format, cell.border = FIN, fmt, BOX
        if label in ("Benefit realisation factor", "Net margin per year on a funded loan", "Discount rate (annual)"):
            cell.fill = FILL_IN
    wi.cell(row=i, column=5, value=note).font = FS
    ref[label] = i
R = lambda label, col="B": f"Inputs!${col}${ref[label]}"

# ---------------------------------------------------------------- Cash flows
wc = wb.create_sheet("Cash_Flows")
wc["A1"], wc["A2"] = "Five-Year Cash Flows by Option (€)", "All cells are formulas on the Inputs sheet."
wc["A1"].font, wc["A2"].font = FT, FS
wc.sheet_view.showGridLines = False
wc.column_dimensions["A"].width = 44
for j in range(2, 9):
    wc.column_dimensions[chr(64 + j)].width = 14
YEARS = ["Year 0", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
blocks = {}
row = 4
for opt, col in [("A. Do nothing", "B"), ("B. Quick wins", "C"), ("C. Digital + quick wins", "D")]:
    wc.cell(row=row, column=1, value=opt).font = FB
    for j, h in enumerate(YEARS + ["Total"], start=2):
        cell = wc.cell(row=row, column=j, value=h)
        cell.font, cell.fill, cell.border = FH, FILL_H, BOX
    lines = {}
    defs = [
        ("Revenue benefit: extra loans funded", lambda y:
            "0" if y == 0 else
            f"{R('Extra funded loans vs today (modelled)', col)}*{R('Benefit realisation factor')}*{R('Year ' + str(y), col)}"
            f"*{R('Average funded loan (€)')}*{R('Net margin per year on a funded loan')}*{R('Average loan life (years)')}"),
        ("Efficiency: underwriter re-keying saved", lambda y:
            "0" if y == 0 else
            f"{R('Credit assessments', col)}*{R('Minutes re-keying financials per assessment', col)}/60"
            f"*{R('Reduction in re-keying', col)}*{R('Underwriter cost per hour (€)')}*{R('Year ' + str(y), col)}"),
        ("Efficiency: RM document chasing saved", lambda y:
            "0" if y == 0 else
            f"{R('Document requests (today)', col)}*{R('Reduction in document requests', col)}"
            f"*{R('Minutes of RM time per document request', col)}/60*{R('Relationship manager cost per hour (€)')}"
            f"*{R('Year ' + str(y), col)}"),
        ("One-off costs", lambda y:
            f"-{R('One-off cost: year 0', col)}" if y == 0 else (f"-{R('One-off cost: year 1', col)}" if y == 1 else "0")),
        ("Run costs", lambda y: "0" if y == 0 else f"-{R('Annual run cost (from year 1)', col)}"),
    ]
    for name, fn in defs:
        row += 1
        wc.cell(row=row, column=1, value=name).font = F
        for y in range(6):
            cell = wc.cell(row=row, column=2 + y, value="=" + fn(y))
            cell.number_format, cell.border, cell.font = EUR, BOX, F
        t = wc.cell(row=row, column=8, value=f"=SUM(B{row}:G{row})")
        t.number_format, t.border, t.font = EUR, BOX, FB
        lines[name] = row
    row += 1
    first, last = lines["Revenue benefit: extra loans funded"], lines["Run costs"]
    wc.cell(row=row, column=1, value="Net cash flow").font = FB
    for j in range(2, 9):
        L = chr(64 + j)
        cell = wc.cell(row=row, column=j, value=f"=SUM({L}{first}:{L}{last})")
        cell.number_format, cell.border, cell.font = EUR, BOX, FB
    lines["net"] = row
    row += 1
    wc.cell(row=row, column=1, value="Cumulative cash flow").font = F
    for y in range(6):
        L = chr(66 + y)
        f = f"={L}{lines['net']}" if y == 0 else f"={chr(65 + y)}{row}+{L}{lines['net']}"
        cell = wc.cell(row=row, column=2 + y, value=f)
        cell.number_format, cell.border, cell.font = EUR, BOX, F
    lines["cum"] = row
    wc.conditional_formatting.add(f"B{row}:G{row}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color="C00000")))
    blocks[opt] = lines
    row += 3

# ---------------------------------------------------------------- Summary
ws = wb.create_sheet("Summary", 0)
ws["A1"], ws["A2"] = "Options Appraisal Summary", "5-year view. NPV at the Inputs discount rate. All values are formulas."
ws["A1"].font, ws["A2"].font = FT, FS
ws.sheet_view.showGridLines = False
for c, w in zip("ABCDEFGH", (30, 16, 12, 16, 16, 18, 18, 16)):
    ws.column_dimensions[c].width = w
heads = ["Option", "NPV (€)", "IRR", "Payback (years)", "Year 0 investment (€)", "5-yr benefits (€)", "5-yr costs (€)", "Extra loans / yr"]
for j, h in enumerate(heads, 1):
    cell = ws.cell(row=4, column=j, value=h)
    cell.font, cell.fill, cell.border = FH, FILL_H, BOX
    cell.alignment = Alignment(wrap_text=True, horizontal="center")
rate = R("Discount rate (annual)")
for i, (opt, col) in enumerate([("A. Do nothing", "B"), ("B. Quick wins", "C"), ("C. Digital + quick wins", "D")]):
    r = 5 + i
    L = blocks[opt]
    n, cum = L["net"], L["cum"]
    rev, eff1, eff2 = L["Revenue benefit: extra loans funded"], L["Efficiency: underwriter re-keying saved"], L["Efficiency: RM document chasing saved"]
    vals = [opt,
            f"=Cash_Flows!B{n}+NPV({rate},Cash_Flows!C{n}:G{n})",
            f'=IF(Cash_Flows!H{n}=0,"n/a",IFERROR(IRR(Cash_Flows!B{n}:G{n}),"n/a"))',
            # payback: first year where cumulative turns positive, interpolated
            f'=IF(Cash_Flows!H{n}=0,"n/a",IFERROR(IF(Cash_Flows!B{cum}>=0,0,'
            f'MATCH(TRUE,INDEX(Cash_Flows!B{cum}:G{cum}>=0,0),0)-2'
            f'-INDEX(Cash_Flows!B{cum}:G{cum},MATCH(TRUE,INDEX(Cash_Flows!B{cum}:G{cum}>=0,0),0)-1)'
            f'/INDEX(Cash_Flows!B{n}:G{n},MATCH(TRUE,INDEX(Cash_Flows!B{cum}:G{cum}>=0,0),0))),"> 5 years"))',
            f"=-Cash_Flows!B{n}",
            f"=Cash_Flows!H{rev}+Cash_Flows!H{eff1}+Cash_Flows!H{eff2}",
            f"=-(Cash_Flows!H{L['One-off costs']}+Cash_Flows!H{L['Run costs']})",
            f"={R('Extra funded loans vs today (modelled)', col)}*{R('Benefit realisation factor')}"]
    fmts = [None, EUR, PCT, "0.0", EUR, EUR, EUR, "#,##0"]
    for j, (v, fm) in enumerate(zip(vals, fmts), 1):
        cell = ws.cell(row=r, column=j, value=v)
        cell.font, cell.border = (FB if j == 1 else F), BOX
        if fm:
            cell.number_format = fm
ws["A9"] = "Recommended option (highest NPV)"
ws["A9"].font = FB
ws["B9"] = "=INDEX(A5:A7,MATCH(MAX(B5:B7),B5:B7,0))"
ws["B9"].font = Font(name=FONT, size=12, bold=True, color="008000")
ws["A10"] = ("Recommendation: approve Option B now (fast payback, low risk) and Option C as phase 2, "
             "funded from Option B's benefits, if the NPV holds at the sensitivity below.")
ws["A10"].font = FS

# Sensitivity on realisation factor
ws["A13"] = "Sensitivity: NPV (€) by benefit realisation factor"
ws["A13"].font = FB
for j, h in enumerate(["Option", "25%", "50%", "75%", "100%"], 1):
    cell = ws.cell(row=14, column=j, value=h if j == 1 else float(h.strip("%")) / 100)
    cell.font, cell.fill, cell.border = FH, FILL_H, BOX
    if j > 1:
        cell.number_format = "0%"
for i, opt in enumerate(["B. Quick wins", "C. Digital + quick wins"]):
    r = 15 + i
    L = blocks[opt]
    rev = L["Revenue benefit: extra loans funded"]
    ws.cell(row=r, column=1, value=opt).font = F
    for j in range(2, 6):
        col = chr(64 + j)
        # NPV(r) = NPV(all flows) + NPV(revenue row) x (r / current realisation - 1)
        f = (f"=B{6 + i}+(Cash_Flows!B{rev}+NPV({rate},Cash_Flows!C{rev}:G{rev}))"
             f"*({col}$14/{R('Benefit realisation factor')}-1)")
        cell = ws.cell(row=r, column=j, value=f)
        cell.number_format, cell.border, cell.font = EUR, BOX, F
ws.conditional_formatting.add("B15:E16", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color="C00000")))

ch = BarChart()
ch.title, ch.style, ch.height, ch.width = "NPV by option (€)", 2, 7, 14
ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=7), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=7))
ch.legend = None
ch.series[0].graphicalProperties.solidFill = "2A78D6"
ch.y_axis.number_format = '€#,##0,,"M"'
ws.add_chart(ch, "A19")

wb.calculation.fullCalcOnLoad = True
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"Business case written: {OUT}")
